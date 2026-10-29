"""Tableaux des PDF vers Excel : reperage des tableaux (traits ou colonnes alignees), lecture des cellules,
conversion des nombres a la francaise, ecriture .xlsx (une feuille par tableau, ou tout dans une seule)."""
import re

import pymupdf as fitz

_NB = re.compile(r"^[-+]?\(?\s*\d{1,3}(?:[  . ]?\d{3})*(?:,\d+)?\s*\)?\s*(?:€|EUR|%)?$|^[-+]?\d+(?:[.,]\d+)?\s*(?:€|EUR|%)?$")
_DATE = re.compile(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})$")


def valeur_cellule(t):
    """Texte d'une cellule -> nombre (1 234,56 € -> 1234.56 ; (12,00) -> -12), date, ou le texte tel quel."""
    if t is None:
        return None
    s = " ".join(str(t).split())
    if not s:
        return None
    if _NB.match(s):
        neg = s.startswith("(") and s.endswith(")") or s.startswith("-")
        pct = s.endswith("%")
        x = re.sub(r"[€%()+\-\s   ]|EUR", "", s)
        if "," in x:
            x = x.replace(".", "").replace(",", ".")
        elif x.count(".") > 1 or re.match(r"^\d{1,3}\.\d{3}$", x):
            x = x.replace(".", "")
        try:
            v = float(x)
        except ValueError:
            return s
        v = -v if neg else v
        v = v / 100 if pct else v
        return int(v) if v == int(v) and "," not in s and "." not in s and not pct else v
    m = _DATE.match(s)
    if m:
        import datetime
        j, mo, a = map(int, m.groups())
        a = a + 2000 if a < 100 else a
        try:
            return datetime.date(a, mo, j)
        except ValueError:
            return s
    return s


def trouver(doc, pages=None, strategie="auto"):
    """Tableaux du document : [{page, rect, lignes: [[texte]], entete: bool}].
    strategie : "traits" (tableaux encadres), "texte" (colonnes alignees sans traits), "auto" (traits, sinon texte)."""
    out = []
    for n in (pages if pages is not None else range(len(doc))):
        page = doc[n]
        essais = ["lines"] if strategie == "traits" else ["text"] if strategie == "texte" else ["lines", "text"]
        for s in essais:
            try:
                tf = page.find_tables(strategy=s)
            except Exception:
                continue
            tabs = [t for t in tf.tables if t.row_count >= 2 and t.col_count >= 2]
            if s == "text":                       # sans traits : on ecarte le texte courant pris pour un tableau
                tabs = [t for t in tabs if _remplissage(t) >= 0.45]
            if tabs:
                for t in tabs:
                    lignes, ys = grille(page, t, avec_y=True)
                    lignes = _rogner(lignes, ys) if s == "text" else [l for l in lignes if any(l)]
                    nc = max((len(l) for l in lignes), default=0)
                    garde = [j for j in range(nc) if any(j < len(l) and l[j] for l in lignes)]
                    lignes = [[l[j] if j < len(l) else "" for j in garde] for l in lignes]
                    if len(lignes) >= 2:
                        out.append({"page": n, "rect": fitz.Rect(t.bbox), "lignes": lignes,
                                    "entete": _a_entete(lignes), "strategie": s})
                break
    return out


def _bornes(t):
    """Limites des colonnes et des lignes du tableau detecte."""
    xs, ys = set(), set()
    for r in t.rows:
        ys.add(round(r.bbox[1], 1))
        ys.add(round(r.bbox[3], 1))
        for c in r.cells:
            if c is not None:
                xs.add(round(c[0], 1))
                xs.add(round(c[2], 1))
    return _fusion(sorted(xs)), _fusion(sorted(ys))


def _fusion(v, tol=2.0):
    out = []
    for x in v:
        if not out or x - out[-1] > tol:
            out.append(x)
    return out


def _couloirs(mots, x0, x1, ys, nb_lignes_mini=2, largeur=11.0):
    """Dans une colonne, bandes verticales vides de tout mot (« couloirs »), avec des mots des deux cotes sur au
    moins deux lignes : la colonne en cache deux (un trait vertical manquant). Renvoie les x de coupe."""
    if not mots:
        return []
    occupe = sorted((max(x0, m[0]), min(x1, m[2])) for m in mots)
    coupes, fin = [], x0
    libres = []
    for a, b in occupe:
        if a - fin >= largeur and fin > x0 + 1:
            libres.append((fin, a))
        fin = max(fin, b)
    for a, b in libres:
        milieu = (a + b) / 2
        lignes = {}
        for m in mots:
            cy = (m[1] + m[3]) / 2
            k = next((i for i in range(len(ys) - 1) if ys[i] <= cy < ys[i + 1]), round(cy))
            lignes.setdefault(k, set()).add("g" if m[2] <= milieu else "d")
        if sum(1 for v in lignes.values() if v == {"g", "d"}) >= nb_lignes_mini:
            coupes.append(milieu)
    return coupes


def grille(page, t, avec_y=False):
    """Cellules du tableau, relues mot a mot (espaces gardes, nombres comme « 130,00 » d'un seul tenant) ;
    une colonne qui en contient deux sans trait de separation est coupee."""
    xs, ys = _bornes(t)
    if len(xs) < 2 or len(ys) < 2:
        l_ = [[(c or "").strip() for c in l] for l in t.extract()]
        return (l_, list(range(len(l_)))) if avec_y else l_
    bbox = fitz.Rect(xs[0], ys[0], xs[-1], ys[-1])
    mots = [m for m in page.get_text("words", clip=bbox + (-1, -1, 1, 1))]
    nouvelles = []
    for a, b in zip(xs, xs[1:]):
        dedans = [m for m in mots if a <= (m[0] + m[2]) / 2 < b]
        nouvelles += _couloirs(dedans, a, b, ys)
    xs = sorted(xs + nouvelles)
    cases = [[[] for _ in range(len(xs) - 1)] for _ in range(len(ys) - 1)]
    for m in mots:
        cx, cy = (m[0] + m[2]) / 2, (m[1] + m[3]) / 2
        i = next((k for k in range(len(ys) - 1) if ys[k] <= cy < ys[k + 1]), None)
        j = next((k for k in range(len(xs) - 1) if xs[k] <= cx < xs[k + 1]), None)
        if i is not None and j is not None:
            cases[i][j].append(m)
    out = []
    for l in cases:
        rangee = []
        for c in l:
            c.sort(key=lambda m: (round(m[1] / 3), m[0]))
            rangee.append(" ".join(m[4] for m in c))
        out.append(rangee)
    # colonnes entierement vides (bords) retirees
    garde = [j for j in range(len(xs) - 1) if any(l[j] for l in out)]
    out = [[l[j] for j in garde] for l in out]
    if avec_y:
        return out, [(a + b) / 2 for a, b in zip(ys, ys[1:])]
    return out


def _rogner(lignes, ys):
    """Sans traits, le titre au-dessus (ou le texte en dessous) est pris dans le tableau : on retire les lignes
    separees du reste par un grand blanc."""
    paires = [(l, y) for l, y in zip(lignes, ys) if any(l)]
    if len(paires) < 4:
        return [l for l, _y in paires]
    ecarts = sorted(b[1] - a[1] for a, b in zip(paires, paires[1:]))
    med = ecarts[len(ecarts) // 2]
    debut, fin = 0, len(paires)
    while debut < fin - 2 and paires[debut + 1][1] - paires[debut][1] > 1.8 * med:
        debut += 1
    while fin - 1 > debut + 2 and paires[fin - 1][1] - paires[fin - 2][1] > 1.8 * med:
        fin -= 1
    return [l for l, _y in paires[debut:fin]]


def _remplissage(t):
    l = [x for x in t.extract() if any(c and c.strip() for c in x)]
    cases = sum(len(x) for x in l) or 1
    return sum(1 for x in l for c in x if c and c.strip()) / cases


def _a_entete(lignes):
    """Premiere ligne = titres de colonnes : que du texte, alors que la suite contient des nombres."""
    if len(lignes) < 2:
        return False
    tete = [valeur_cellule(c) for c in lignes[0]]
    if any(isinstance(v, (int, float)) for v in tete):
        return False
    suite = [valeur_cellule(c) for l in lignes[1:] for c in l]
    return any(isinstance(v, (int, float)) for v in suite) or len(set(lignes[0])) == len(lignes[0])


def assembler(tableaux):
    """Tableaux qui se suivent sur plusieurs pages avec les memes colonnes (un releve de 6 pages) : un seul.
    La ligne d'en-tete repetee en haut de chaque page n'est gardee qu'une fois."""
    out = []
    for t in tableaux:
        p = out[-1] if out else None
        if p and t["page"] == p["page_fin"] + 1 and len(t["lignes"][0]) == len(p["lignes"][0]):
            lignes = t["lignes"][1:] if p["entete"] and t["lignes"][0] == p["lignes"][0] else t["lignes"]
            p["lignes"] += lignes
            p["page_fin"] = t["page"]
            p["cadres"].append((t["page"], t.get("rect")))     # le cadre de la suite, sur sa page
            continue
        out.append(dict(t, lignes=[list(l) for l in t["lignes"]], page_fin=t["page"],
                        cadres=[(t["page"], t.get("rect"))]))
    return out


def ecrire_xlsx(chemin, feuilles, convertir=True):
    """feuilles : [(nom, lignes)] ; premiere ligne en gras si c'est un en-tete. Nombres et dates convertis."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    vus = set()
    for nom, lignes in feuilles:
        nom = re.sub(r"[\[\]:*?/\\]", "-", str(nom))[:31] or "Feuille"
        base, k = nom, 2
        while nom.lower() in vus:
            nom = f"{base[:28]} {k}"
            k += 1
        vus.add(nom.lower())
        ws = wb.create_sheet(nom)
        largeurs = {}
        for i, l in enumerate(lignes, 1):
            for j, c in enumerate(l, 1):
                v = valeur_cellule(c) if (convertir and isinstance(c, str)) else c
                cell = ws.cell(row=i, column=j, value=v)
                if hasattr(v, "year"):
                    cell.number_format = "DD/MM/YYYY"
                elif isinstance(v, float):
                    cell.number_format = "#,##0.00"
                if isinstance(c, str) and "\n" in c:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                largeurs[j] = max(largeurs.get(j, 6), min(60, max((len(x) for x in str(c or "").split("\n")),
                                                                    default=0) + 2))
        if lignes and _a_entete([[str(c) if c is not None else "" for c in l] for l in lignes]) or \
                (lignes and all(isinstance(c, str) for c in lignes[0])):
            for c in ws[1]:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="DDE7F3")
            ws.freeze_panes = "A2"
        for j, w in largeurs.items():
            ws.column_dimensions[get_column_letter(j)].width = w
    if not wb.sheetnames:
        wb.create_sheet("Feuille")
    wb.save(chemin)
    return chemin
