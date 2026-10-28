"""Outil Mesure : echelle (rapport, lue sur le plan, etalonnage), distance, perimetre, surface, liste, Excel.
Lancement : QT_QPA_PLATFORM=offscreen python tests/test_mesure.py"""
import os, sys, tempfile
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ["PHILPDF_LANGUE"] = "fr"
os.environ["APPDATA"] = tempfile.mkdtemp()
os.environ["HOME"] = os.environ["APPDATA"]
import annotations  # noqa: F401
import pymupdf as fitz
from PySide6.QtWidgets import QApplication, QDialog, QInputDialog
from PySide6.QtCore import Qt
app = QApplication.instance() or QApplication([])
from app.config import Parametres
from app.document import OngletDocument
import app.dialogues_mesure as DM
from moteur import mesure as M, tableaux as TB

ok = []


def verif(c, quoi):
    ok.append((bool(c), quoi))
    print(("OK    " if c else "ECHEC ") + quoi)


T = tempfile.mkdtemp()
# moteur
e = M.depuis_rapport(200)
verif(abs(M.valeurs("distance", [(0, 0), (72 / 2.54, 0)], e)["longueur"] - 2) < 1e-9, "1:200 : 1 cm sur le papier = 2 m")
verif(abs(M.valeurs("surface", [(0, 0), (300, 0), (300, 200), (0, 200)], e)["aire"] - 298.6944) < 0.01, "surface (lacet)")
e2 = M.depuis_equivalence(1, 1, "m")
verif(abs(M.rapport(e2) - 100) < 0.01, "1 cm = 1 m équivaut à 1:100")
e3 = M.depuis_etalon(100, 5, "m")
verif(abs(M.valeurs("distance", [(0, 0), (200, 0)], e3)["longueur"] - 10) < 1e-9, "étalonnage sur une longueur connue")
verif(M.texte_valeur("surface", {"aire": 12345.678, "longueur": 0}, "m").startswith("12 345,68 m²"), "nombre à la française et ha")
d = fitz.open()
p = d.new_page(width=842, height=595)
p.insert_text((50, 50), "Echelle d'origine : 1/1000   Echelle d'édition : 1/500")
p.draw_rect(fitz.Rect(100, 100, 400, 300))
verif(M.lire_texte(d, 0)[0] == 500, "échelle lue sur le plan (édition avant origine)")
src = os.path.join(T, "plan.pdf")
d.save(src)
# interface
o = OngletDocument(Parametres(), src)
o.resize(1200, 800)
o.show()
DM.DialogueEchelle.exec = lambda self: (self.valider(), QDialog.Accepted)[-1]
o.changer_mode("mesure")
verif(o.bandeau_mesure.isVisible() and o.echelle_page(0)["texte"] == "1:500", "mode Mesure : échelle demandée, lue sur le plan")
pw = o.page_widgets[0]
o.choisir_genre_mesure("surface")
for q in [(100, 100), (400, 100), (400, 300), (100, 300)]:
    o.clic_mesure(pw, fitz.Point(q), Qt.NoModifier)
o.finir_mesure()
o.choisir_genre_mesure("distance")
o.clic_mesure(pw, fitz.Point(100, 350), Qt.NoModifier)
o.clic_mesure(pw, fitz.Point(401.5, 351), Qt.ShiftModifier)          # Maj : horizontal
ms = M.lister(o.doc)
verif(len(ms) == 2 and ms[0]["genre"] == "surface" and "1 866,78 m²" in ms[0]["texte"], f"surface posée ({ms[0]['texte'] if ms else ''})")
verif(ms[1]["genre"] == "distance" and ms[1]["texte"].startswith("53,1"), f"distance posée, Maj = horizontale ({ms[1]['texte'] if len(ms) > 1 else ''})")
# aimant : un clic a 3 pt du coin s'accroche au coin
o.choisir_genre_mesure("perimetre")
o.clic_mesure(pw, fitz.Point(102, 98), Qt.NoModifier)
verif(o.mesure["points"][-1] == (100.0, 100.0), "aimant sur les coins du plan")
o.clic_mesure(pw, fitz.Point(250, 100), Qt.NoModifier)
o.annuler_point_mesure()
verif(len(o.mesure["points"]) == 1, "retour arrière : dernier point retiré")
o.abandonner_mesure()
# etalonnage
o.echelles.clear()
DM.DialogueEchelle.exec = lambda self: (self.r_etalon.setChecked(True), self.valider(), QDialog.Accepted)[-1]
o.regler_echelle(0)
DM.demander_longueur = lambda parent, u: 30.0
o.clic_mesure(pw, fitz.Point(100, 400), Qt.NoModifier)
o.clic_mesure(pw, fitz.Point(400, 400), Qt.NoModifier)
verif(abs(o.echelle_page(0)["c"] - 0.1) < 1e-9, "étalonnage : 300 pt = 30 m")
# Acrobat : echelle gardee dans le PDF
vp = M.lire_vp(o.doc, 0)
verif(vp and vp.get("philpdf") and abs(vp["c"] - 0.1) < 1e-6, "échelle gardée dans le PDF (/VP /Measure)")
# renommer, liste, Excel
M.renommer(o.doc, 0, ms[0]["ident"], "Parcelle 12")
verif(M.lister(o.doc)[0]["texte"].startswith("Parcelle 12 : "), "mesure renommée")
dl = DM.DialogueMesures(o)
verif(dl.table.rowCount() == 2 and "Total surface" in dl.totaux.text(), "liste des mesures et totaux")
x = os.path.join(T, "mesures.xlsx")
TB.ecrire_xlsx(x, [("Mesures", dl.lignes())])
from openpyxl import load_workbook
ws = load_workbook(x).active
verif(ws["B2"].value == "Parcelle 12" and isinstance(ws["D2"].value, float), "export Excel des mesures (nombres)")
# enregistrement : les mesures restent
out = os.path.join(T, "plan_mesure.pdf")
o.doc.save(out)
verif(len(M.lister(fitz.open(out))) == 2, "mesures gardées dans le fichier")
M.supprimer(o.doc, 0, ms[1]["ident"])
verif(len(M.lister(o.doc)) == 1 and len(list(o.doc[0].annots())) == 2, "suppression : mesure et étiquette")
o.changer_mode(None)
verif(not o.bandeau_mesure.isVisible(), "sortie du mode Mesure")
n_ok = sum(1 for c, _q in ok if c)
print(f"\n{n_ok}/{len(ok)} vérifications de mesure réussies")
sys.exit(0 if n_ok == len(ok) else 1)
