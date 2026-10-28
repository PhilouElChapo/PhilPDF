"""Tableaux vers Excel : reperage (traits, colonnes alignees), relecture des cellules, nombres et dates a la
francaise, tableau sur plusieurs pages, apercu, ecriture .xlsx. Lancement : QT_QPA_PLATFORM=offscreen python tests/test_excel.py"""
import datetime, os, sys, tempfile
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ["PHILPDF_LANGUE"] = "fr"
os.environ["APPDATA"] = tempfile.mkdtemp()
os.environ["HOME"] = os.environ["APPDATA"]
import annotations  # noqa: F401
import pymupdf as fitz
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox
app = QApplication.instance() or QApplication([])
from moteur import tableaux as TB

ok = []


def verif(c, quoi):
    ok.append((bool(c), quoi))
    print(("OK    " if c else "ECHEC ") + quoi)


T = tempfile.mkdtemp()
verif(TB.valeur_cellule("1 234,56 €") == 1234.56 and TB.valeur_cellule("(12,00)") == -12.0, "montants français")
verif(TB.valeur_cellule("06/10/2026") == datetime.date(2026, 10, 6) and TB.valeur_cellule("Total HT") == "Total HT",
      "dates et texte")
verif(TB.valeur_cellule("1.250") == 1250 and TB.valeur_cellule("20%") == 0.2, "milliers à point, pourcentage")
# facture encadree (traits) dont deux colonnes n'ont pas de trait entre elles
d = fitz.open()
p = d.new_page()
x = [50, 330, 420, 480, 545]
ys = [100, 120, 140, 160]
for y in ys:
    p.draw_line((x[0], y), (x[-1], y))
for xx in x:
    p.draw_line((xx, ys[0]), (xx, ys[-1]))
for i, ligne in enumerate([["Désignation", "Quantité", "Prix unitaire", "TVA", "Montant HT"],
                           ["Dépannage portail", "1,00", "130,00", "20%", "130,00"],
                           ["Déplacement", "2,00", "25,00", "20%", "50,00"]]):
    y = ys[i] + 14
    p.insert_text((54, y), ligne[0], fontsize=9)
    p.insert_text((280, y), ligne[1], fontsize=9)        # meme case que la designation, sans trait
    for k, v in enumerate(ligne[2:]):
        p.insert_text((x[k + 1] + 4, y), v, fontsize=9)
# releve sans traits sur deux pages, en-tete repete
for pg in range(2):
    q = d.new_page()
    q.insert_text((50, 60), "Relevé de compte — page " + str(pg + 1), fontsize=12)
    for i in range(16):
        y = 100 + i * 18
        lig = (["Date", "Libellé", "Débit", "Crédit"] if i == 0 else
               [f"{i:02d}/0{pg + 1}/2026", f"Opération numéro {pg * 20 + i}", f"{i * 3},50" if i % 4 else "", "" if i % 4 else f"{i * 10},00"])
        for xx, v in zip((50, 130, 400, 480), lig):
            q.insert_text((xx, y), v, fontsize=9)
src = os.path.join(T, "tableaux.pdf")
d.save(src)
d = fitz.open(src)
t = TB.trouver(d, pages=[0])
verif(t and t[0]["lignes"][0] == ["Désignation", "Quantité", "Prix unitaire", "TVA", "Montant HT"],
      f"facture : colonne sans trait coupée ({t[0]['lignes'][0] if t else ''})")
verif(t and t[0]["lignes"][1][0] == "Dépannage portail" and t[0]["entete"], "facture : mots avec espaces, en-tête")
r = TB.trouver(d, pages=[1, 2], strategie="texte")
verif(len(r) == 2 and len(r[0]["lignes"][0]) == 4, f"relevé sans traits : colonnes alignées ({[len(x['lignes'][0]) for x in r]})")
a = TB.assembler(r)
verif(len(a) == 1 and len(a[0]["lignes"]) == 31 and a[0]["page_fin"] == 2,
      f"relevé sur deux pages réuni, en-tête gardé une fois ({len(a[0]['lignes']) if a else 0} lignes)")
# dialogue
from app.config import Parametres
from app.document import OngletDocument
import app.dialogues_excel as DX
o = OngletDocument(Parametres(), src)
dl = DX.DialogueExcel(o)
dl.resize(1180, 760)
dl.show()
app.processEvents()
verif(dl.liste.count() >= 2, f"aperçu : tableaux listés ({dl.liste.count()})")
dl.liste.setCurrentRow(0)
app.processEvents()
verif(dl.grille.horizontalHeaderItem(0).text() == "Désignation" and dl.grille.item(0, 4).text() == "130,00",
      "aperçu : cellules telles qu'elles iront dans Excel")
dl.grab().save(os.path.join(T, "excel.png"))
dl.aller(0)
nb = dl.liste.count()
dl.tableau_trace(fitz.Rect(45, 95, 550, 165))
verif(dl.liste.count() == nb + 1, "tableau tracé à la main")
dl.retirer()
f = os.path.join(T, "sortie.xlsx")
QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (f, ""))
QMessageBox.question = staticmethod(lambda *a, **k: QMessageBox.No)
dl.exporter()
from openpyxl import load_workbook
wb = load_workbook(f)
ws = wb.worksheets[0]
verif(ws["A1"].value == "Désignation" and ws["E2"].value == 130 and ws["A1"].font.bold, "Excel : en-tête en gras, nombres")
ws2 = wb.worksheets[1]
verif(isinstance(ws2["A2"].value, datetime.datetime) and ws2["C2"].value == 3.5 and ws2["D5"].value == 40, "Excel : dates et montants du relevé")
print("capture :", os.path.join(T, "excel.png"))
n_ok = sum(1 for c, _q in ok if c)
print(f"\n{n_ok}/{len(ok)} vérifications Excel réussies")
sys.exit(0 if n_ok == len(ok) else 1)
